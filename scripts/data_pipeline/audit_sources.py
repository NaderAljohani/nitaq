from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
FILES = [
    ROOT / "data/raw/core/gastat_warehousing_logistics_2023_en.xlsx",
    ROOT / "data/raw/core/gastat_warehousing_logistics_2024_en.xlsx",
    ROOT / "data/raw/core/gastat_land_transport_2023_en.xlsx",
    ROOT / "data/raw/core/gastat_road_transport_2024_ar.xlsx",
]
PATTERN = re.compile(
    r"table\s*(?:1|2|5|8|9|13)\b|delivery|warehouse|logistic|freight|postal|region|"
    r"جدول|طلبات\s*التوصيل|المستودعات|المراكز\s*اللوجستية|نقل\s*البضائع|البريد|المنطقة",
    re.IGNORECASE,
)


def text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


result = []
for path in FILES:
    wb = load_workbook(path, read_only=True, data_only=False)
    workbook = {"file": path.name, "sheets": []}
    for ws in wb.worksheets:
        matches = []
        sample_rows = []
        for row_idx, row in enumerate(ws.iter_rows(), start=1):
            values = [text(c.value) for c in row]
            if any(values) and len(sample_rows) < 12:
                while values and not values[-1]:
                    values.pop()
                sample_rows.append({"row": row_idx, "values": values[:20]})
            for col_idx, value in enumerate(values, start=1):
                if value and PATTERN.search(value):
                    matches.append({"cell": f"{ws.cell(row_idx, col_idx).coordinate}", "text": value[:300]})
                    if len(matches) >= 80:
                        break
            if len(matches) >= 80:
                break
        workbook["sheets"].append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "matches": matches,
            "sample_rows": sample_rows,
        })
    wb.close()
    result.append(workbook)

out = ROOT / "outputs/build/temp/source_audit.json"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({
    "output": str(out),
    "workbooks": [
        {"file": w["file"], "sheets": [{"name": s["name"], "rows": s["max_row"], "cols": s["max_col"], "matches": len(s["matches"])} for s in w["sheets"]]}
        for w in result
    ],
}, ensure_ascii=False, indent=2))
