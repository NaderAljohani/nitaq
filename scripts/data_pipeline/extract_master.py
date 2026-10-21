from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CORE = ROOT / "data/raw/core"
PROCESSED = ROOT / "data/processed"
TMP = ROOT / "outputs/build/temp"

WAREHOUSE_2023 = CORE / "gastat_warehousing_logistics_2023_en.xlsx"
WAREHOUSE_2024 = CORE / "gastat_warehousing_logistics_2024_en.xlsx"
POPULATION_KAPSARC = CORE / "saudi-arabia-population-by-administrative-region-nationality-and-sex.xlsx"
POPULATION_DERIVED = PROCESSED / "population_by_region_2023_2024.csv"

OUT_CSV = PROCESSED / "mizan_region_master_2024.csv"
OUT_DICTIONARY = PROCESSED / "mizan_region_master_2024_dictionary.csv"
OUT_JSON = TMP / "mizan_region_master_2024.json"
OUT_QA = ROOT / "outputs/build/mizan_region_master_2024_validation.json"

REGIONS = [
    (1, "Riyadh", "الرياض"),
    (2, "Makkah", "مكة المكرمة"),
    (3, "Madinah", "المدينة المنورة"),
    (4, "Qassim", "القصيم"),
    (5, "Eastern Region", "المنطقة الشرقية"),
    (6, "Aseer", "عسير"),
    (7, "Tabuk", "تبوك"),
    (8, "Hail", "حائل"),
    (9, "Northern Borders", "الحدود الشمالية"),
    (10, "Jazan", "جازان"),
    (11, "Najran", "نجران"),
    (12, "Al-Baha", "الباحة"),
    (13, "Al-Jouf", "الجوف"),
]
REGION_ORDER = {en: i for i, en, _ in REGIONS}
REGION_AR = {en: ar for _, en, ar in REGIONS}

ALIASES = {
    "Riyadh": "Riyadh",
    "Makkah": "Makkah",
    "Madinah": "Madinah",
    "Al-Madinah": "Madinah",
    "Qassim": "Qassim",
    "Al-Qassim": "Qassim",
    "Eastern Region": "Eastern Region",
    "Eastern region": "Eastern Region",
    "Aseer": "Aseer",
    "Asir": "Aseer",
    "Tabuk": "Tabuk",
    "Hail": "Hail",
    "Northern Borders": "Northern Borders",
    "Jazan": "Jazan",
    "Najran": "Najran",
    "Al-Baha": "Al-Baha",
    "Al-Bahah": "Al-Baha",
    "Al-Jouf": "Al-Jouf",
    "Al-Jawf": "Al-Jouf",
}


def canonical(value: str) -> str:
    key = " ".join(str(value).split())
    if key not in ALIASES:
        raise ValueError(f"Unknown region name: {value!r}")
    return ALIASES[key]


def growth(current, prior):
    if prior in (None, 0):
        return None
    return round((current / prior) - 1, 6)


def ratio(numerator, denominator, scale=1):
    if denominator in (None, 0):
        return None
    return round(numerator / denominator * scale, 6)


def read_population():
    wb = load_workbook(POPULATION_KAPSARC, read_only=True, data_only=True)
    ws = wb.active
    direct = {2023: {}, 2024: {}}
    national = {}
    for row in ws.iter_rows(min_col=1, max_col=4, values_only=True):
        date, region, gender, population = row
        if str(date) not in {"2023", "2024"} or gender != "Total":
            continue
        year = int(date)
        if region == "Total":
            national[year] = int(population)
        else:
            direct[year][canonical(region)] = int(population)
    wb.close()

    derived = {}
    with POPULATION_DERIVED.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            region = canonical(row["region_en"])
            derived[region] = {
                2023: int(row["population_2023"]),
                2024: int(row["population_2024"]),
            }
    return direct, national, derived


def read_warehouses():
    wb24 = load_workbook(WAREHOUSE_2024, read_only=True, data_only=True)
    ws24 = wb24["5"]
    data = {}
    for row in ws24.iter_rows(min_row=8, max_row=20, min_col=2, max_col=6, values_only=True):
        region, count23, area23, count24, area24 = row
        key = canonical(region)
        data[key] = {
            "warehouse_count_2023": int(count23),
            "warehouse_area_sqm_2023": round(float(area23), 2),
            "warehouse_count_2024": int(count24),
            "warehouse_area_sqm_2024": round(float(area24), 2),
        }
    source_totals = {
        "warehouse_count_2023": int(ws24["C21"].value),
        "warehouse_area_sqm_2023": round(float(ws24["D21"].value), 2),
        "warehouse_count_2024": int(ws24["E21"].value),
        "warehouse_area_sqm_2024": round(float(ws24["F21"].value), 2),
    }
    wb24.close()

    wb23 = load_workbook(WAREHOUSE_2023, read_only=True, data_only=True)
    ws23 = wb23["5"]
    crosscheck = {}
    for row in ws23.iter_rows(min_row=7, max_row=19, min_col=2, max_col=4, values_only=True):
        region, count, area = row
        key = canonical(region)
        crosscheck[key] = (int(count), round(float(area), 2))
    wb23.close()
    return data, source_totals, crosscheck


def read_licenses():
    result = {region: {} for _, region, _ in REGIONS}
    wb23 = load_workbook(WAREHOUSE_2023, read_only=True, data_only=True)
    ws23 = wb23["8"]
    for region, value in ws23.iter_rows(min_row=7, max_row=19, min_col=2, max_col=3, values_only=True):
        result[canonical(region)]["freight_post_licenses_2023"] = int(value)
    total23 = int(ws23["C20"].value)
    wb23.close()

    wb24 = load_workbook(WAREHOUSE_2024, read_only=True, data_only=True)
    ws24 = wb24["9"]
    for region, value in ws24.iter_rows(min_row=6, max_row=18, min_col=2, max_col=3, values_only=True):
        result[canonical(region)]["freight_post_licenses_2024"] = int(value)
    total24 = int(ws24["C19"].value)
    wb24.close()
    return result, {"freight_post_licenses_2023": total23, "freight_post_licenses_2024": total24}


def read_delivery():
    wb = load_workbook(WAREHOUSE_2024, read_only=True, data_only=True)
    ws = wb["13"]
    data = {}
    for row in ws.iter_rows(min_row=8, max_row=20, min_col=2, max_col=7, values_only=True):
        region, q1, q2, q3, q4, total = row
        key = canonical(region)
        data[key] = {
            "delivery_q1_2024": int(q1),
            "delivery_q2_2024": int(q2),
            "delivery_q3_2024": int(q3),
            "delivery_q4_2024": int(q4),
            "delivery_total_2024_source": int(total),
        }
    source_totals = {
        "delivery_q1_2024": int(ws["C21"].value),
        "delivery_q2_2024": int(ws["D21"].value),
        "delivery_q3_2024": int(ws["E21"].value),
        "delivery_q4_2024": int(ws["F21"].value),
        "delivery_total_2024_source": int(ws["G21"].value),
    }
    wb.close()
    return data, source_totals


def read_centers():
    wb = load_workbook(WAREHOUSE_2024, read_only=True, data_only=True)
    ws = wb["2"]
    data = defaultdict(lambda: {
        "activated_logistics_centers_2023": 0,
        "activated_logistics_centers_2024": 0,
        "activated_logistics_centers_area_sqm_2023": 0,
        "activated_logistics_centers_area_sqm_2024": 0,
    })
    current_region = None
    for row in ws.iter_rows(min_row=8, max_row=30, min_col=2, max_col=16, values_only=True):
        region = row[0]
        if region:
            current_region = canonical(region)
        if current_region is None:
            raise ValueError("Center row encountered before region label")
        area23_thousand = row[7]
        status23 = row[8]
        area24_thousand = row[13]
        status24 = row[14]
        if status23 == "Activated":
            data[current_region]["activated_logistics_centers_2023"] += 1
            data[current_region]["activated_logistics_centers_area_sqm_2023"] += int(round(float(area23_thousand) * 1000))
        if status24 == "Activated":
            data[current_region]["activated_logistics_centers_2024"] += 1
            data[current_region]["activated_logistics_centers_area_sqm_2024"] += int(round(float(area24_thousand) * 1000))
    national_count23 = int(wb["1"]["F7"].value)
    national_count24 = int(wb["1"]["G7"].value)
    wb.close()

    wb23 = load_workbook(WAREHOUSE_2023, read_only=True, data_only=True)
    ws23 = wb23["2"]
    check = defaultdict(lambda: {"count": 0, "area_sqm": 0})
    for row in ws23.iter_rows(min_row=8, max_row=29, min_col=2, max_col=5, values_only=True):
        _, region, _, area_thousand = row
        key = canonical(region)
        check[key]["count"] += 1
        check[key]["area_sqm"] += int(round(float(area_thousand) * 1000))
    wb23.close()
    return data, {"activated_logistics_centers_2023": national_count23, "activated_logistics_centers_2024": national_count24}, check


population, population_national, population_derived = read_population()
warehouses, warehouse_source_totals, warehouse_crosscheck = read_warehouses()
licenses, license_source_totals = read_licenses()
delivery, delivery_source_totals = read_delivery()
centers, center_source_totals, center_crosscheck = read_centers()

rows = []
for region_id, region_en, region_ar in REGIONS:
    pop23 = population[2023][region_en]
    pop24 = population[2024][region_en]
    wh = warehouses[region_en]
    lic = licenses[region_en]
    dlv = delivery[region_en]
    ctr = centers[region_en]
    delivery_calc = sum(dlv[f"delivery_q{quarter}_2024"] for quarter in range(1, 5))
    row = {
        "region_id": region_id,
        "region_en": region_en,
        "region_ar": region_ar,
        "population_2023": pop23,
        "population_2024": pop24,
        **dlv,
        "delivery_total_2024_calculated": delivery_calc,
        "delivery_total_matches_quarters": delivery_calc == dlv["delivery_total_2024_source"],
        **wh,
        **lic,
        **ctr,
        "population_growth_pct": growth(pop24, pop23),
        "delivery_q4_vs_q1_growth_pct": growth(dlv["delivery_q4_2024"], dlv["delivery_q1_2024"]),
        "delivery_requests_per_1000_population_2024": ratio(dlv["delivery_total_2024_source"], pop24, 1000),
        "warehouse_count_growth_pct": growth(wh["warehouse_count_2024"], wh["warehouse_count_2023"]),
        "warehouse_area_growth_pct": growth(wh["warehouse_area_sqm_2024"], wh["warehouse_area_sqm_2023"]),
        "freight_post_licenses_growth_pct": growth(lic["freight_post_licenses_2024"], lic["freight_post_licenses_2023"]),
        "activated_logistics_centers_change": ctr["activated_logistics_centers_2024"] - ctr["activated_logistics_centers_2023"],
        "activated_logistics_centers_growth_pct": growth(ctr["activated_logistics_centers_2024"], ctr["activated_logistics_centers_2023"]),
        "activated_logistics_centers_area_growth_pct": growth(ctr["activated_logistics_centers_area_sqm_2024"], ctr["activated_logistics_centers_area_sqm_2023"]),
        "warehouse_area_sqm_per_capita_2023": ratio(wh["warehouse_area_sqm_2023"], pop23),
        "warehouse_area_sqm_per_capita_2024": ratio(wh["warehouse_area_sqm_2024"], pop24),
        "warehouse_count_per_100k_population_2024": ratio(wh["warehouse_count_2024"], pop24, 100000),
        "freight_post_licenses_per_100k_population_2024": ratio(lic["freight_post_licenses_2024"], pop24, 100000),
        "activated_logistics_centers_per_million_population_2024": ratio(ctr["activated_logistics_centers_2024"], pop24, 1000000),
        "delivery_requests_per_warehouse_sqm_2024": ratio(dlv["delivery_total_2024_source"], wh["warehouse_area_sqm_2024"]),
        "delivery_requests_per_freight_post_license_2024": ratio(dlv["delivery_total_2024_source"], lic["freight_post_licenses_2024"]),
        "logistics_center_presence_2024": int(ctr["activated_logistics_centers_2024"] > 0),
        "population_source_ref": "KAPSARC population-by-region dataset, Sheet1, Date 2023/2024, Gender=Total; GASTAT road-density recovery retained as validation only",
        "delivery_source_ref": "GASTAT Warehousing and Logistics Statistics 2024, Table 13",
        "warehouse_source_ref": "GASTAT Warehousing and Logistics Statistics 2024, Table 5; 2023 values cross-checked to 2023 edition Table 5",
        "license_source_ref": "GASTAT Warehousing and Logistics Statistics 2023 Table 8 and 2024 Table 9; underlying source TGA",
        "logistics_centers_source_ref": "GASTAT Warehousing and Logistics Statistics 2024, Tables 1-2; underlying source MOTLS Logisti",
        "quality_notes_ar": "السكان مأخوذون مباشرة من KAPSARC؛ 2024 مطابق للاستعادة السابقة إقليميًا، و2023 يطابقها وطنيًا فقط بسبب تقريب كثافة الطرق المنشورة. مجموع أرباع التوصيل مطابق للإجمالي؛ قيم مستودعات 2023 مطابقة بين الإصدارين؛ المراكز المفعلة مجمعة من حالة كل مركز.",
        "limitations_ar": "السكان تقديرات سنوية وليست تعدادًا؛ لا تستخدم الاستعادة من كثافة الطرق كبديل إقليمي لعام 2023؛ المستودعات تراخيص ومساحات مرخصة وليست إشغالًا أو سعة تشغيلية؛ الطلبات تخص الأنشطة المرخصة؛ صفر المراكز يعني عدم ظهور مركز مفعّل وفق تعريف الجدول ولا يعني غياب كل أصل لوجستي.",
    }
    rows.append(row)

COLUMNS = list(rows[0].keys())

dictionary = {
    "region_id": ("معرف ترتيب المنطقة", "identifier", "static", "Canonical region list", "معرف داخلي ثابت"),
    "region_en": ("اسم المنطقة الإنجليزية الموحد", "text", "static", "Canonical region list", "تم توحيد تهجئات المصادر"),
    "region_ar": ("اسم المنطقة العربية الموحد", "text", "static", "Canonical region list", "13 منطقة إدارية"),
    "population_2023": ("سكان المنطقة لعام 2023", "person", "2023", "KAPSARC Sheet1, Date=2023, Gender=Total", "قيمة مباشرة؛ تختلف إقليميًا عن الاستعادة من كثافة الطرق مع تطابق المجموع الوطني"),
    "population_2024": ("سكان المنطقة لعام 2024", "person", "2024", "KAPSARC Sheet1, Date=2024, Gender=Total", "قيمة مباشرة ومطابقة للاستعادة السابقة إقليميًا"),
    "delivery_q1_2024": ("طلبات التوصيل المنفذة في الربع الأول", "request", "2024 Q1", "GASTAT Warehousing and Logistics 2024 Table 13", "نشاط توصيل مرخص"),
    "delivery_q2_2024": ("طلبات التوصيل المنفذة في الربع الثاني", "request", "2024 Q2", "GASTAT Warehousing and Logistics 2024 Table 13", "نشاط توصيل مرخص"),
    "delivery_q3_2024": ("طلبات التوصيل المنفذة في الربع الثالث", "request", "2024 Q3", "GASTAT Warehousing and Logistics 2024 Table 13", "نشاط توصيل مرخص"),
    "delivery_q4_2024": ("طلبات التوصيل المنفذة في الربع الرابع", "request", "2024 Q4", "GASTAT Warehousing and Logistics 2024 Table 13", "نشاط توصيل مرخص"),
    "delivery_total_2024_source": ("إجمالي الطلبات كما نشره المصدر", "request", "2024", "GASTAT Warehousing and Logistics 2024 Table 13", "يُراجع مقابل مجموع الأرباع"),
    "delivery_total_2024_calculated": ("مجموع الأرباع المحسوب", "request", "2024", "Derived: Q1+Q2+Q3+Q4", "يجب أن يطابق إجمالي المصدر"),
    "delivery_total_matches_quarters": ("اختبار مطابقة الإجمالي للأرباع", "boolean", "2024", "Derived validation", "True لكل المناطق"),
    "warehouse_count_2023": ("عدد المستودعات التجارية المرخصة", "licensed warehouse", "2023", "GASTAT Warehousing and Logistics 2024 Table 5", "مطابق لإصدار 2023 Table 5"),
    "warehouse_area_sqm_2023": ("مساحة المستودعات التجارية المرخصة", "square meter", "2023", "GASTAT Warehousing and Logistics 2024 Table 5", "ليست إشغالًا أو سعة تشغيلية"),
    "warehouse_count_2024": ("عدد المستودعات التجارية المرخصة", "licensed warehouse", "2024", "GASTAT Warehousing and Logistics 2024 Table 5", "قيمة ترخيصية"),
    "warehouse_area_sqm_2024": ("مساحة المستودعات التجارية المرخصة", "square meter", "2024", "GASTAT Warehousing and Logistics 2024 Table 5", "ليست إشغالًا أو سعة تشغيلية"),
    "freight_post_licenses_2023": ("تراخيص نقل البضائع والخدمات البريدية السارية", "active license", "end-2023", "GASTAT Warehousing and Logistics 2023 Table 8", "تجميع لأنشطة متعددة حسب حاشية المصدر"),
    "freight_post_licenses_2024": ("تراخيص نقل البضائع والخدمات البريدية النشطة", "active license", "end-2024", "GASTAT Warehousing and Logistics 2024 Table 9", "تجميع لأنشطة متعددة حسب حاشية المصدر"),
    "activated_logistics_centers_2023": ("عدد المراكز اللوجستية المفعلة", "activated center", "2023", "GASTAT Warehousing and Logistics 2024 Tables 1-2", "وفق تعريف المصدر للمراكز المفعلة"),
    "activated_logistics_centers_2024": ("عدد المراكز اللوجستية المفعلة", "activated center", "2024", "GASTAT Warehousing and Logistics 2024 Tables 1-2", "وفق تعريف المصدر للمراكز المفعلة"),
    "activated_logistics_centers_area_sqm_2023": ("مساحة المراكز المفعلة", "square meter", "2023", "GASTAT Warehousing and Logistics 2024 Table 2", "مجموع مساحات المراكز ذات الحالة Activated"),
    "activated_logistics_centers_area_sqm_2024": ("مساحة المراكز المفعلة", "square meter", "2024", "GASTAT Warehousing and Logistics 2024 Table 2", "مجموع مساحات المراكز ذات الحالة Activated"),
    "population_growth_pct": ("نمو السكان", "decimal rate", "2023-2024", "Derived", "(2024/2023)-1"),
    "delivery_q4_vs_q1_growth_pct": ("تغير طلبات التوصيل داخل 2024", "decimal rate", "Q1-Q4 2024", "Derived", "ليس نموًا سنويًا؛ (Q4/Q1)-1"),
    "delivery_requests_per_1000_population_2024": ("طلبات التوصيل لكل ألف نسمة", "request per 1,000 persons", "2024", "Derived", "إجمالي الطلبات / سكان 2024 × 1000"),
    "warehouse_count_growth_pct": ("نمو عدد المستودعات", "decimal rate", "2023-2024", "Derived", "(2024/2023)-1"),
    "warehouse_area_growth_pct": ("نمو مساحة المستودعات", "decimal rate", "2023-2024", "Derived", "(2024/2023)-1"),
    "freight_post_licenses_growth_pct": ("نمو تراخيص النقل والبريد", "decimal rate", "2023-2024", "Derived", "(2024/2023)-1"),
    "activated_logistics_centers_change": ("التغير العددي في المراكز المفعلة", "center", "2023-2024", "Derived", "عدد 2024 ناقص عدد 2023"),
    "activated_logistics_centers_growth_pct": ("نمو عدد المراكز المفعلة", "decimal rate", "2023-2024", "Derived", "فارغ عندما أساس 2023 يساوي صفرًا"),
    "activated_logistics_centers_area_growth_pct": ("نمو مساحة المراكز المفعلة", "decimal rate", "2023-2024", "Derived", "فارغ عندما أساس 2023 يساوي صفرًا"),
    "warehouse_area_sqm_per_capita_2023": ("مساحة مستودعات مرخصة للفرد", "square meter per person", "2023", "Derived", "المساحة / السكان"),
    "warehouse_area_sqm_per_capita_2024": ("مساحة مستودعات مرخصة للفرد", "square meter per person", "2024", "Derived", "المساحة / السكان"),
    "warehouse_count_per_100k_population_2024": ("عدد المستودعات لكل 100 ألف نسمة", "warehouse per 100,000 persons", "2024", "Derived", "العدد / السكان × 100000"),
    "freight_post_licenses_per_100k_population_2024": ("التراخيص لكل 100 ألف نسمة", "license per 100,000 persons", "2024", "Derived", "العدد / السكان × 100000"),
    "activated_logistics_centers_per_million_population_2024": ("المراكز المفعلة لكل مليون نسمة", "center per 1,000,000 persons", "2024", "Derived", "العدد / السكان × 1000000"),
    "delivery_requests_per_warehouse_sqm_2024": ("ضغط الطلبات على مساحة المستودعات المرخصة", "request per square meter", "2024", "Derived", "إجمالي الطلبات / المساحة المرخصة"),
    "delivery_requests_per_freight_post_license_2024": ("طلبات التوصيل لكل ترخيص نقل/بريد", "request per license", "2024", "Derived", "إجمالي الطلبات / التراخيص"),
    "logistics_center_presence_2024": ("وجود مركز لوجستي مفعّل", "binary", "2024", "Derived", "1 إذا كان العدد أكبر من صفر؛ وإلا 0"),
    "population_source_ref": ("مرجع السكان", "text", "2023-2024", "Metadata", "مرجع ومصدر التحقق"),
    "delivery_source_ref": ("مرجع التوصيل", "text", "2024", "Metadata", "اسم الملف والجدول"),
    "warehouse_source_ref": ("مرجع المستودعات", "text", "2023-2024", "Metadata", "اسم الملف والجدول"),
    "license_source_ref": ("مرجع التراخيص", "text", "2023-2024", "Metadata", "اسم الملف والجدول"),
    "logistics_centers_source_ref": ("مرجع المراكز", "text", "2023-2024", "Metadata", "اسم الملف والجداول"),
    "quality_notes_ar": ("ملاحظات الجودة", "text", "all", "Metadata", "نتائج التحقق الأساسية"),
    "limitations_ar": ("قيود الاستخدام", "text", "all", "Metadata", "لا تُستخدم الدرجة كإثبات نقص فعلي"),
}

missing_dictionary = [column for column in COLUMNS if column not in dictionary]
if missing_dictionary:
    raise AssertionError(f"Dictionary entries missing: {missing_dictionary}")

PROCESSED.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
OUT_QA.parent.mkdir(parents=True, exist_ok=True)

with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=COLUMNS)
    writer.writeheader()
    writer.writerows(rows)

dictionary_columns = ["metric", "label_ar", "unit", "period", "source", "definition_or_limitation"]
with OUT_DICTIONARY.open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=dictionary_columns)
    writer.writeheader()
    for metric in COLUMNS:
        label, unit, period, source, note = dictionary[metric]
        writer.writerow({
            "metric": metric,
            "label_ar": label,
            "unit": unit,
            "period": period,
            "source": source,
            "definition_or_limitation": note,
        })

calculated_totals = {}
sum_metrics = [
    "population_2023", "population_2024",
    "delivery_q1_2024", "delivery_q2_2024", "delivery_q3_2024", "delivery_q4_2024", "delivery_total_2024_source",
    "warehouse_count_2023", "warehouse_area_sqm_2023", "warehouse_count_2024", "warehouse_area_sqm_2024",
    "freight_post_licenses_2023", "freight_post_licenses_2024",
    "activated_logistics_centers_2023", "activated_logistics_centers_2024",
    "activated_logistics_centers_area_sqm_2023", "activated_logistics_centers_area_sqm_2024",
]
for metric in sum_metrics:
    calculated_totals[metric] = round(sum(row[metric] for row in rows), 2)

expected_totals = {
    "population_2023": population_national[2023],
    "population_2024": population_national[2024],
    **delivery_source_totals,
    **warehouse_source_totals,
    **license_source_totals,
    **center_source_totals,
    "activated_logistics_centers_area_sqm_2023": 34190000,
    "activated_logistics_centers_area_sqm_2024": 34561000,
}

total_checks = {}
for metric, expected in expected_totals.items():
    calculated = calculated_totals[metric]
    tolerance = 0.01 if isinstance(expected, float) else 0
    total_checks[metric] = {
        "calculated": calculated,
        "expected_source": expected,
        "difference": round(calculated - expected, 6),
        "pass": abs(calculated - expected) <= tolerance,
    }

warehouse_crosscheck_results = {}
for region_en in REGION_ORDER:
    expected_count, expected_area = warehouse_crosscheck[region_en]
    current = warehouses[region_en]
    warehouse_crosscheck_results[region_en] = {
        "count_match": current["warehouse_count_2023"] == expected_count,
        "area_match": math.isclose(current["warehouse_area_sqm_2023"], expected_area, abs_tol=0.01),
    }

population_crosscheck_results = {}
for region_en in REGION_ORDER:
    population_crosscheck_results[region_en] = {
        "population_2023_match": population[2023][region_en] == population_derived[region_en][2023],
        "population_2023_difference_direct_minus_road_density_recovery": population[2023][region_en] - population_derived[region_en][2023],
        "population_2024_match": population[2024][region_en] == population_derived[region_en][2024],
        "population_2024_difference_direct_minus_road_density_recovery": population[2024][region_en] - population_derived[region_en][2024],
    }

population_crosscheck_summary = {
    "population_2023_regional_exact_matches": sum(item["population_2023_match"] for item in population_crosscheck_results.values()),
    "population_2024_regional_exact_matches": sum(item["population_2024_match"] for item in population_crosscheck_results.values()),
    "population_2023_direct_national_total": population_national[2023],
    "population_2023_road_density_recovery_national_total": sum(item[2023] for item in population_derived.values()),
    "population_2024_direct_national_total": population_national[2024],
    "population_2024_road_density_recovery_national_total": sum(item[2024] for item in population_derived.values()),
}
population_crosscheck_summary["population_2023_national_total_match"] = (
    population_crosscheck_summary["population_2023_direct_national_total"]
    == population_crosscheck_summary["population_2023_road_density_recovery_national_total"]
)
population_crosscheck_summary["population_2024_national_total_match"] = (
    population_crosscheck_summary["population_2024_direct_national_total"]
    == population_crosscheck_summary["population_2024_road_density_recovery_national_total"]
)

center_crosscheck_results = {}
for region_en in REGION_ORDER:
    current = centers[region_en]
    prior = center_crosscheck[region_en]
    center_crosscheck_results[region_en] = {
        "count_match": current["activated_logistics_centers_2023"] == prior["count"],
        "area_match": current["activated_logistics_centers_area_sqm_2023"] == prior["area_sqm"],
    }

qa = {
    "dataset": str(OUT_CSV.relative_to(ROOT)),
    "row_count": len(rows),
    "expected_region_count": 13,
    "unique_region_en_count": len({row["region_en"] for row in rows}),
    "unique_region_ar_count": len({row["region_ar"] for row in rows}),
    "duplicate_region_keys": len(rows) - len({row["region_en"] for row in rows}),
    "missing_raw_values": {
        column: sum(row[column] is None or row[column] == "" for row in rows)
        for column in COLUMNS[:22]
    },
    "all_delivery_quarter_sums_match": all(row["delivery_total_matches_quarters"] for row in rows),
    "source_total_checks": total_checks,
    "population_kapsarc_vs_gastat_derived": population_crosscheck_results,
    "population_crosscheck_summary": population_crosscheck_summary,
    "warehouse_2023_crosscheck_between_editions": warehouse_crosscheck_results,
    "centers_2023_crosscheck_between_editions": center_crosscheck_results,
    "raw_files_preserved": [
        str(WAREHOUSE_2023.relative_to(ROOT)),
        str(WAREHOUSE_2024.relative_to(ROOT)),
        str(POPULATION_KAPSARC.relative_to(ROOT)),
        str(POPULATION_DERIVED.relative_to(ROOT)),
    ],
}
qa["all_checks_pass"] = (
    qa["row_count"] == 13
    and qa["unique_region_en_count"] == 13
    and qa["unique_region_ar_count"] == 13
    and qa["duplicate_region_keys"] == 0
    and all(value == 0 for value in qa["missing_raw_values"].values())
    and qa["all_delivery_quarter_sums_match"]
    and all(item["pass"] for item in total_checks.values())
    and population_crosscheck_summary["population_2023_national_total_match"]
    and population_crosscheck_summary["population_2024_national_total_match"]
    and population_crosscheck_summary["population_2024_regional_exact_matches"] == 13
    and all(all(item.values()) for item in warehouse_crosscheck_results.values())
    and all(all(item.values()) for item in center_crosscheck_results.values())
)

OUT_JSON.write_text(json.dumps({"rows": rows, "dictionary": dictionary, "qa": qa}, ensure_ascii=False, indent=2), encoding="utf-8")
OUT_QA.write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding="utf-8")

print(json.dumps({
    "rows": len(rows),
    "columns": len(COLUMNS),
    "all_checks_pass": qa["all_checks_pass"],
    "output_csv": str(OUT_CSV),
    "output_dictionary": str(OUT_DICTIONARY),
    "output_qa": str(OUT_QA),
    "totals": calculated_totals,
}, ensure_ascii=False, indent=2))
