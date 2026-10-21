from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SELECTIONS = {
    "gastat_warehousing_logistics_2023_en.xlsx": ["1", "2", "5", "8"],
    "gastat_warehousing_logistics_2024_en.xlsx": ["1", "2", "5", "9", "13"],
}


def compact(value):
    if value is None:
        return None
    if isinstance(value, str):
        return " ".join(value.split())
    return value


result = {}
for filename, sheets in SELECTIONS.items():
    path = ROOT / "data/raw/core" / filename
    wb = load_workbook(path, read_only=True, data_only=True)
    result[filename] = {}
    for sheet_name in sheets:
        ws = wb[sheet_name]
        rows = []
        for row_number, row in enumerate(ws.iter_rows(), start=1):
            values = [compact(cell.value) for cell in row]
            if any(value not in (None, "") for value in values):
                while values and values[-1] in (None, ""):
                    values.pop()
                rows.append({"row": row_number, "values": values})
            if row_number >= 120:
                break
        result[filename][sheet_name] = rows
    wb.close()

out = ROOT / "outputs/build/temp/selected_tables.json"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(str(out))
